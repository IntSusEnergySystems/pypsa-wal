#!/usr/bin/env python3
"""Build the one-slide weather-year deliverable for the cabinet deck.

Wraps the figure written by ``compare_weather_years.py --bare`` in a 16:9 slide:
title, one-line subtitle, the figure, and a single takeaway strip. Deliberately
sparse — the figure carries the argument, the text only frames it.

The slide is self-contained (no master, no template), so it is meant to be
*copied into* the real deck rather than presented as-is. All wording lives in
the ``TEXT`` dict below: switch ``--lang`` to ``fr`` for the French version, or
edit the dict for any other phrasing.

Requires ``python-pptx``, which is **not** in the ``pypsa-eur`` environment. On
this machine it lives in ``pptx_latex_extractor``::

    ~/progs/miniconda/envs/pypsa-eur/bin/python \\
        scripts/walloon_scripts/compare_weather_years.py --bare [--lang fr]
    ~/progs/miniconda/envs/pptx_latex_extractor/bin/python \\
        scripts/walloon_scripts/make_weather_year_slide.py [--lang fr]

Usage::

    python scripts/walloon_scripts/make_weather_year_slide.py \\
        --lang fr      # picks weather_year_BEWAL_2050_bare_fr.png automatically
"""

from __future__ import annotations

import argparse
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

INK = RGBColor(0x0B, 0x0B, 0x0B)
INK2 = RGBColor(0x52, 0x51, 0x4E)
ACCENT = RGBColor(0xEB, 0x68, 0x34)  # the 2010 series colour
SURFACE = RGBColor(0xFC, 0xFC, 0xFB)
TINT = RGBColor(0xF7, 0xEF, 0xE9)
FONT = "Calibri"

TEXT = {
    "en": {
        "title": "The weather year moves the supply, not the demand",
        "subtitle": "Wallonia, 2050 — TIMES sets the annual demand; the weather "
                    "year sets the renewable resource and when it arrives",
        "lead": "Why 2010 — ",
        "takeaway": "the year's weakest renewable month is also its coldest. "
                    "The system answers with +11% wind and +82% battery storage.",
        "notes": (
            "Why the weather year matters even though PyPSA does not set the demand.\n\n"
            "1. Demand. TIMES fixes the annual energy of every Walloon demand — "
            "electricity, heat, transport, industry. Checked load by load: identical "
            "to 1e-11 %, 105.7 TWh in both runs. The weather year cannot move it.\n\n"
            "2. Resource. The weather year drives the wind, PV and run-of-river "
            "profiles. On an identical fleet, 2010 yields 27.9 TWh of wind + PV "
            "against 28.4 TWh in 2013: wind -5.1 %, solar +2.5 %. On the annual "
            "average 2010 is barely poorer.\n\n"
            "3. Timing — the actual reason 2010 is the stress case. The same weather "
            "also redistributes the fixed TIMES heat envelope across the 8 760 hours: "
            "13.6 % of the demand moves to other hours. In December 2010 the two "
            "effects compound — renewable output 35 % below December 2013 while heat "
            "demand runs 47 % above it. Over the winter as a whole: renewables "
            "-13.9 %, heat +7.6 %.\n\n"
            "4. Response. The optimiser sizes the system for that December: "
            "+10.6 % onshore wind (6 500 vs 5 879 MW), +5.8 % rooftop PV, "
            "+81.7 % battery storage (6.5 vs 3.6 GWh).\n\n"
            "Caveat worth stating if asked: because TIMES fixes the annual heat "
            "energy, a cold year shifts heat within the year but never raises the "
            "yearly total. A real 2010 would have consumed more heat overall, so "
            "what is shown is the timing component only — the 2010 case is "
            "conservative on that count.\n\n"
            "Source: scen_central (2010) vs scen_central_2013, cabinet batch of "
            "2026-09-13, 1 h resolution. Figure: "
            "scripts/walloon_scripts/compare_weather_years.py"
        ),
    },
    "fr": {
        "title": "L'année météo déplace la production, pas la demande",
        "subtitle": "Wallonie, 2050 — TIMES fixe la demande annuelle ; l'année "
                    "météo fixe la ressource renouvelable et son calendrier",
        "lead": "Pourquoi 2010 — ",
        "takeaway": "le mois le moins renouvelable de l'année est aussi le plus "
                    "froid. Le système répond par +11 % d'éolien et +82 % de batteries.",
        "notes": (
            "Pourquoi l'année météo compte alors que PyPSA ne fixe pas la demande.\n\n"
            "1. Demande. TIMES fixe l'énergie annuelle de chaque demande wallonne — "
            "électricité, chaleur, transport, industrie. Vérifié charge par charge : "
            "identique à 1e-11 % près, 105,7 TWh dans les deux runs.\n\n"
            "2. Ressource. L'année météo détermine les profils éolien, PV et fil de "
            "l'eau. À parc identique, 2010 produit 27,9 TWh d'éolien + PV contre "
            "28,4 TWh en 2013 : éolien -5,1 %, solaire +2,5 %.\n\n"
            "3. Calendrier — la vraie raison. La météo redistribue aussi l'enveloppe "
            "de chaleur fixée par TIMES sur les 8 760 heures : 13,6 % de la demande "
            "change d'heure. En décembre 2010 les deux effets se cumulent : "
            "production renouvelable 35 % sous décembre 2013, demande de chaleur "
            "47 % au-dessus. Sur l'hiver : renouvelables -13,9 %, chaleur +7,6 %.\n\n"
            "4. Réponse. L'optimiseur dimensionne le système sur ce décembre : "
            "+10,6 % d'éolien terrestre, +5,8 % de PV en toiture, +81,7 % de "
            "batteries (6,5 contre 3,6 GWh).\n\n"
            "Nuance à mentionner si la question vient : TIMES fixant l'énergie "
            "annuelle de chaleur, une année froide déplace la chaleur dans l'année "
            "mais n'augmente jamais le total annuel. Seul l'effet de calendrier est "
            "représenté — le cas 2010 est donc conservateur sur ce point.\n\n"
            "Source : scen_central (2010) vs scen_central_2013, batch cabinet du "
            "13-09-2026, résolution horaire."
        ),
    },
}

SLIDE_W, SLIDE_H = 13.333, 7.5
MARGIN = 0.55
BODY_W = SLIDE_W - 2 * MARGIN


def textbox(slide, left, top, width, height, *, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    return tf


def run(paragraph, text, *, size, color, bold=False, font=FONT):
    r = paragraph.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.color.rgb = color
    r.font.bold = bold
    r.font.name = font
    return r


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--lang", choices=sorted(TEXT), default="en")
    ap.add_argument("--figure", type=Path, default=None,
                    help="default: the --bare figure matching --lang")
    ap.add_argument("--out", type=Path, default=None)
    args = ap.parse_args()
    t = TEXT[args.lang]
    tag = "" if args.lang == "en" else f"_{args.lang}"
    if args.figure is None:
        args.figure = Path(f"docs/figures/weather_year_BEWAL_2050_bare{tag}.png")
    if args.out is None:
        args.out = Path(f"docs/figures/weather_year_slide{tag}.pptx")

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(SLIDE_W), Inches(SLIDE_H)
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank

    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0,
                                Inches(SLIDE_W), Inches(SLIDE_H))
    bg.fill.solid()
    bg.fill.fore_color.rgb = SURFACE
    bg.line.fill.background()
    bg.shadow.inherit = False

    tf = textbox(slide, MARGIN, 0.30, BODY_W, 0.60)
    run(tf.paragraphs[0], t["title"], size=30, color=INK, bold=True)

    tf = textbox(slide, MARGIN, 0.98, BODY_W, 0.40)
    run(tf.paragraphs[0], t["subtitle"], size=13.5, color=INK2)

    # Scale by width only, so the height follows the figure's own aspect and the
    # layout below adapts instead of silently stretching it.
    img_top = 1.44
    pic = slide.shapes.add_picture(str(args.figure), Inches(MARGIN),
                                   Inches(img_top), width=Inches(BODY_W))
    img_h = Emu(pic.height).inches

    strip_top = img_top + img_h + 0.18
    strip = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(MARGIN),
                                   Inches(strip_top), Inches(BODY_W), Inches(0.52))
    strip.fill.solid()
    strip.fill.fore_color.rgb = TINT
    strip.line.fill.background()
    strip.shadow.inherit = False
    strip.adjustments[0] = 0.28
    stf = strip.text_frame
    stf.word_wrap = True
    stf.vertical_anchor = MSO_ANCHOR.MIDDLE
    stf.margin_left = Inches(0.20)
    stf.margin_right = Inches(0.20)
    stf.margin_top = stf.margin_bottom = 0
    p = stf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    run(p, t["lead"], size=13, color=ACCENT, bold=True)
    run(p, t["takeaway"], size=13, color=INK)

    slide.notes_slide.notes_text_frame.text = t["notes"]

    args.out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(args.out))
    bottom_margin = SLIDE_H - (strip_top + 0.52)
    print(f"written: {args.out}  ({args.lang}, figure {args.figure.name})")
    print(f"  image {BODY_W:.2f} x {img_h:.2f} in, bottom margin "
          f"{bottom_margin:.2f} in" + ("" if bottom_margin >= 0.4 else "  <-- TIGHT"))


if __name__ == "__main__":
    main()
